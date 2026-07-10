import asyncio
import base64
import os
import tempfile
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from fastapi import UploadFile
from faster_whisper import WhisperModel
from pypdf import PdfReader


ContentPart = dict[str, Any]

_whisper_model: WhisperModel | None = None
_whisper_lock = asyncio.Lock()


def extract_pdf_text(
    data: bytes,
    max_pages: int = 50,
) -> str:
    reader = PdfReader(BytesIO(data))
    parts: list[str] = []

    for index, page in enumerate(reader.pages):
        if index >= max_pages:
            break

        parts.append(page.extract_text() or "")

    text = "\n\n".join(parts).strip()

    if len(text) < 100 and len(reader.pages) >= 5:
        return "[это скан, OCR пока не поддерживается]"

    return text


def extract_docx_text(data: bytes) -> str:
    document = Document(BytesIO(data))
    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def get_whisper_model() -> WhisperModel:
    global _whisper_model

    if _whisper_model is None:
        _whisper_model = WhisperModel(
            "base",
            device="cpu",
            compute_type="int8",
        )

    return _whisper_model


def transcribe_audio_sync(
    audio_bytes: bytes,
    filename: str,
) -> str:
    suffix = Path(filename).suffix or ".ogg"
    temporary_path: str | None = None

    try:
        with tempfile.NamedTemporaryFile(
            suffix=suffix,
            delete=False,
        ) as temporary_file:
            temporary_file.write(audio_bytes)
            temporary_path = temporary_file.name

        model = get_whisper_model()

        segments, _ = model.transcribe(
            temporary_path,
            language="ru",
            beam_size=5,
            vad_filter=True,
            condition_on_previous_text=False,
            initial_prompt=(
                "Это голосовое сообщение на русском языке. "
                "Точно распознавай технические термины: "
                "нейросеть, искусственный интеллект, "
                "Telegram, backend, API, Python."
            ),
        )

        transcript_parts = [
            segment.text.strip()
            for segment in segments
            if segment.text.strip()
        ]

        transcript = " ".join(transcript_parts).strip()

        if not transcript:
            raise ValueError(
                "Не удалось распознать речь в аудиосообщении"
            )

        return transcript

    finally:
        if temporary_path and os.path.exists(temporary_path):
            os.unlink(temporary_path)


async def whisper_transcribe(
    audio_bytes: bytes,
    filename: str,
) -> str:
    async with _whisper_lock:
        return await asyncio.to_thread(
            transcribe_audio_sync,
            audio_bytes,
            filename,
        )


async def media_to_part(
    media: UploadFile,
    llm_client: Any | None = None,
) -> ContentPart:
    del llm_client

    mime = media.content_type or ""
    data = await media.read()

    if mime.startswith("image/"):
        encoded = base64.b64encode(data).decode("ascii")

        return {
            "type": "image_url",
            "image_url": {
                "url": f"data:{mime};base64,{encoded}",
            },
        }

    if mime.startswith("audio/") or mime == "application/ogg":
        transcript = await whisper_transcribe(
            data,
            media.filename or "audio.ogg",
        )

        return {
            "type": "text",
            "text": (
                "[Распознанный текст голосового сообщения]\n"
                f"{transcript}"
            ),
        }

    if mime == "application/pdf":
        text = extract_pdf_text(data)[:3_000]

        return {
            "type": "text",
            "text": f"[Содержимое PDF-документа]\n{text}",
        }

    if mime.endswith("wordprocessingml.document"):
        text = extract_docx_text(data)[:3_000]

        return {
            "type": "text",
            "text": f"[Содержимое DOCX-документа]\n{text}",
        }

    raise ValueError(f"Unsupported media type: {mime}")