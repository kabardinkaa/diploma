from __future__ import annotations

import json
import shutil
import zipfile
from collections import defaultdict
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from reportlab.lib.styles import getSampleStyleSheet

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_PATH = PROJECT_ROOT / "data" / "support_kb.json"
OUTPUT_ROOT = PROJECT_ROOT / "data" / "corporate"
FORMATS = (".md", ".html", ".docx", ".pdf")
FIXED_TIME = datetime(2026, 1, 1, tzinfo=UTC)
FIXED_ZIP_TIME = (2026, 1, 1, 0, 0, 0)


def content_for(item: dict) -> tuple[str, list[str]]:
    title = item["title"]
    paragraphs = [
        item["text"],
        (
            "Материал предназначен для учебного стенда внутренней технической "
            "поддержки контактного центра."
        ),
        (
            "Перед изменением настроек сотрудник фиксирует текст ошибки и время "
            "сбоя. Если описанные шаги не помогли, создаётся заявка с категорией "
            f"«{item['category']}» и приложенными диагностическими данными."
        ),
        (
            "Пароли, коды MFA и другие секреты в заявку, чат и снимки экрана "
            "не добавляются."
        ),
    ]
    return title, paragraphs


def write_markdown(path: Path, title: str, paragraphs: list[str]) -> None:
    body = [f"# {title}", "", *sum(([text, ""] for text in paragraphs), [])]
    path.write_text("\n".join(body).rstrip() + "\n", encoding="utf-8")


def write_html(path: Path, title: str, paragraphs: list[str]) -> None:
    body = "\n".join(f"    <p>{text}</p>" for text in paragraphs)
    html = (
        "<!doctype html>\n<html lang=\"ru\">\n<head>\n"
        "  <meta charset=\"utf-8\">\n"
        f"  <title>{title}</title>\n</head>\n<body>\n"
        f"  <section id=\"guide\">\n    <h1>{title}</h1>\n{body}\n"
        "  </section>\n</body>\n</html>\n"
    )
    path.write_text(html, encoding="utf-8")


def normalize_docx_archive(path: Path) -> None:
    temporary = path.with_suffix(".normalized")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temporary,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    temporary.replace(path)


def write_docx(path: Path, title: str, paragraphs: list[str]) -> None:
    document = Document()
    document.core_properties.author = "Diploma course corpus generator"
    document.core_properties.created = FIXED_TIME
    document.core_properties.modified = FIXED_TIME
    document.add_heading(title, level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(path)
    normalize_docx_archive(path)


def find_unicode_font() -> Path:
    candidates = (
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/dejavu/DejaVuSans.ttf"),
    )
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise RuntimeError("A Cyrillic TrueType font is required to build PDF fixtures")


def write_pdf(path: Path, title: str, paragraphs: list[str]) -> None:
    font_name = "CorpusUnicode"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, str(find_unicode_font())))
    styles = getSampleStyleSheet()
    for style_name in ("Title", "BodyText"):
        styles[style_name].fontName = font_name
    document = SimpleDocTemplate(
        str(path),
        title=title,
        author="Diploma course corpus generator",
        invariant=1,
    )
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    for text in paragraphs:
        story.extend([Paragraph(text, styles["BodyText"]), Spacer(1, 8)])
    document.build(story)


def main() -> None:
    resolved_output = OUTPUT_ROOT.resolve()
    if PROJECT_ROOT.resolve() not in resolved_output.parents:
        raise RuntimeError("Refusing to replace corpus outside the project")
    if OUTPUT_ROOT.exists():
        shutil.rmtree(OUTPUT_ROOT)
    OUTPUT_ROOT.mkdir(parents=True)

    items = json.loads(SOURCE_PATH.read_text(encoding="utf-8"))
    by_category: dict[str, list[dict]] = defaultdict(list)
    for item in items:
        if not item.get("is_archived"):
            by_category[item["category"]].append(item)

    writers = {
        ".md": write_markdown,
        ".html": write_html,
        ".docx": write_docx,
        ".pdf": write_pdf,
    }
    generated = 0
    for category in sorted(by_category):
        category_dir = OUTPUT_ROOT / category
        category_dir.mkdir()
        for index, item in enumerate(by_category[category][:4], start=1):
            extension = FORMATS[generated % len(FORMATS)]
            path = category_dir / f"{category}_guide_{index:02d}_v2026{extension}"
            title, paragraphs = content_for(item)
            writers[extension](path, title, paragraphs)
            generated += 1

    print(f"Generated {generated} documents in {OUTPUT_ROOT}")


if __name__ == "__main__":
    main()
